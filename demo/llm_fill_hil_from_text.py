"""Fill HIL signal columns from the text case workbook with Qwen semantic matching.

The workbook layout used here is:
  - case workbook: text steps in columns F-I, HIL output in columns J-M
  - requirement workbook: SignalName in column C, Name In Chinese in column D

POWERUP and POWERDOWN are deterministic fixed actions. Other Chinese phrases are
matched to requirement rows by the LLM, then rendered as write#/read# commands.
"""

import argparse
import json
import os
import re
import time
import urllib.request
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font

try:
    from demo.embedding_retriever import EmbeddingRetriever
except ModuleNotFoundError:
    from embedding_retriever import EmbeddingRetriever


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_PROVIDER = os.environ.get("LLM_PROVIDER") or os.environ.get("API_PROVIDER", "qwen").lower()
PROVIDER_DEFAULTS = {
    "qwen": {
        "model": os.environ.get("QWEN_MODEL", "qwen-plus"),
        "base_url": os.environ.get("QWEN_BASE_URL", "https://dashscope.aliyuncs.com/compatible-mode/v1"),
        "embedding_model": os.environ.get("QWEN_EMBEDDING_MODEL", "text-embedding-v4"),
        "api_key": os.environ.get("QWEN_API_KEY") or os.environ.get("DASHSCOPE_API_KEY"),
    },
    "openai": {
        "model": os.environ.get("OPENAI_MODEL", "gpt-4.1-mini"),
        "base_url": os.environ.get("OPENAI_BASE_URL", "https://api.openai.com/v1"),
        "embedding_model": os.environ.get("OPENAI_EMBEDDING_MODEL", "text-embedding-3-small"),
        "api_key": os.environ.get("OPENAI_API_KEY"),
    },
}


def provider_defaults(provider):
    provider = ("" if provider is None else str(provider)).strip().lower() or "qwen"
    return PROVIDER_DEFAULTS.get(provider, PROVIDER_DEFAULTS["qwen"])


DEFAULTS = provider_defaults(DEFAULT_PROVIDER)
DEFAULT_MODEL = DEFAULTS["model"]
DEFAULT_BASE_URL = DEFAULTS["base_url"]
DEFAULT_EMBEDDING_MODEL = DEFAULTS["embedding_model"]

SYSTEM_PROMPT = """你是汽车 HIL 测试用例信号匹配助手。
任务：从需求表候选信号中，找出与用例中文短语语义最匹配的 SignalName。

规则：
1. 只返回 JSON，不要解释。
2. 必须从 candidates 中已有的 signalName 选择，禁止编造 candidates 之外的 SignalName。
3. actionRole=write 时优先选择“请求/设置/开关/输入/目标”类信号。
4. actionRole=read 时优先选择“实际/反馈/状态/高度/车速/档位”类信号。
5. 匹配 SignalName 时主要比较 phrase 与 candidates.chineseName 的语义；Description 主要用于解释 value，不要因为 Description 里的相近词改选别的中文名。
6. 必须区分部件层级、系统层级、信号类型和位置语义；不能只因为候选与 phrase 共享“高度/模式/状态”等通用词就匹配。
7. 匹配时要求目标对象一致、实际/目标/请求/反馈/有效/故障等信号类型一致、位置维度一致；如果这些关键语义不同，应降低优先级或返回空。
8. 如果候选存在 FL/FR/RL/RR 等位置后缀，且 phrase 表达多个轮端或多个位置，应返回对应多个位置明确的信号；不要返回有效标志、故障、目标值或控制状态来代替实际测量值。
9. 如果 phrase 包含“指定...”占位词，也要理解它对应的真实信号中文名。
10. 默认只返回 1 个最匹配信号；只有 maxMatches > 1 且 phrase 明确要求多个轮端/多个高度/多个信号时，才返回多个 SignalName。
11. value 必须由你语义理解提取：例如“已挂入P挡”返回 P，“车速=0km/h”返回 0，“激活/开启/有效”按 Description 枚举返回对应值。
12. 当 phrase 有“指定...”时，结合参数上下文中的具体场景值以及 candidates.description 的枚举说明返回 value。
13. value 应该是可直接写入/比较的原始值，例如 0、1、2、0x0、P；不要返回“标准/运动/自适应”等中文含义，除非 Description 明确规定物理值就是中文。
14. 对“初始状态检测”“控制器无故障”“车辆静止”这类没有明确可控/可读信号名的泛化步骤，返回空列表。
15. 信心不足时返回空列表；匹配到信号但无法确定取值时 value 返回空字符串。

返回格式：
{
  "matches": [
    {"signalName": "ExactSignalName", "value": "从用例和Description推断出的值或空", "chineseName": "候选中文名", "confidence": 0.0, "reason": "简短原因"}
  ]
}
"""

EXTRACT_TARGET_PROMPT = """你是汽车 HIL 测试用例语义抽取助手。
任务：从用例短语中抽取它想操作或检查的需求表中文信号名。

规则：
1. 只返回 JSON，不要解释。
2. 返回的 targetChineseNames 应该尽量接近需求表 Name In Chinese 的写法。
3. 不要返回取值本身，例如 P挡、0km/h、标准、运动、自适应不是信号名。
4. 如果短语只是泛化动作，没有明确信号名，返回空列表。
5. phrase 中的“指定...”通常是信号中文名占位，要抽取它代表的信号名。
6. actionRole=write 且短语是“设置X/请求X/输入X”时，targetChineseNames 应该偏向 X请求，不要抽成 X反馈。
7. actionRole=read 且短语是“检测实际X/反馈X/状态X”时，targetChineseNames 应该偏向 X反馈/实际X/状态X。
8. 同时返回 searchQueries，用于在需求表 Name In Chinese 中语义检索；searchQueries 可以包含同义或上位说法，但仍然必须是信号名风格，不要包含取值。
9. 扩展 searchQueries 时必须保留 phrase 中的核心对象、信号类型和位置维度；可以补充同义词、全称、缩写、系统名和位置化表达，但不要泛化到不同对象或不同信号类型。
10. 如果短语涉及 FL/FR/RL/RR、前后左右、四轮或多个位置，searchQueries 应包含位置明确的候选表达。
11. 另外返回 signalProfile，描述当前短语对应的对象层级、信号类型和位置维度，例如 {"object":"suspension","kind":"actual_value","positions":["FL","FR"]}，用于后续召回时过滤明显不一致的候选。

返回格式：
{
  "targetChineseNames": ["中文信号名1", "中文信号名2"],
  "searchQueries": ["检索短语1", "检索短语2"],
  "signalProfile": {
    "object": "object",
    "kind": "actual_value|request|feedback|valid|fault|status|target",
    "positions": ["FL", "FR"]
  }
}
"""


@dataclass(frozen=True)
class RequirementSignal:
    row: int
    signal_name: str
    chinese_name: str
    sender: str
    description: str


def normalize(value):
    return "" if value is None else str(value).replace("\u3000", " ").strip()


def compact_text(value):
    return re.sub(r"\s+", "", normalize(value))


def default_file(exact_name):
    path = PROJECT_ROOT / exact_name
    if path.exists():
        return path
    matches = [p for p in PROJECT_ROOT.glob("*.xlsx") if p.name == exact_name]
    if matches:
        return matches[0]
    raise FileNotFoundError(f"Cannot find {exact_name} in {PROJECT_ROOT}")


def default_requirement_file():
    for path in [
        PROJECT_ROOT / "demo" / "需求1-1_完整版_AI.xlsx",
        PROJECT_ROOT / "demo" / "需求1-1_AI填充.xlsx",
        PROJECT_ROOT / "需求1-1.xlsx",
    ]:
        if path.exists():
            return path
    return default_file("需求1-1.xlsx")


def find_header_col(sheet, *needles, default=None):
    for col in range(1, sheet.max_column + 1):
        header = normalize(sheet.cell(1, col).value).lower()
        if all(needle.lower() in header for needle in needles):
            return col
    return default


def find_hil_columns(sheet):
    for col in range(1, sheet.max_column + 1):
        value = normalize(sheet.cell(1, col).value)
        if "HIL" in value:
            return {
                "initial": col,
                "step": col + 1,
                "action": col + 2,
                "expected": col + 3,
            }
    return {"initial": 10, "step": 11, "action": 12, "expected": 13}


def text_columns_from_hil(hil_columns):
    return {
        "initial": hil_columns["initial"] - 4,
        "step": hil_columns["step"] - 4,
        "action": hil_columns["action"] - 4,
        "expected": hil_columns["expected"] - 4,
    }


def case_ranges(sheet):
    starts = [row for row in range(3, sheet.max_row + 1) if sheet.cell(row, 1).value]
    merged_by_start = {
        merged.min_row: merged.max_row
        for merged in sheet.merged_cells.ranges
        if merged.min_col == 1 and merged.max_col == 1
    }
    ranges = []
    for index, start in enumerate(starts):
        next_start = starts[index + 1] if index + 1 < len(starts) else sheet.max_row + 1
        end = min(merged_by_start.get(start, next_start - 1), next_start - 1)
        ranges.append((start, end))
    return ranges


def split_items(value):
    text = normalize(value)
    if not text or text == "/":
        return []
    text = re.sub(r"(?<!^)(?<![\r\n])\s*(\d+[、.)）])", r"\n\1", text)
    text = re.sub(r"^\s*\d+[、.)）]\s*", "", text, flags=re.M)
    return [item.strip() for item in re.split(r"[;\r\n]+", text) if item.strip()]


def load_requirement_signals(requirement_path):
    workbook = load_workbook(requirement_path, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    signal_col = find_header_col(sheet, "SignalName", default=3)
    chinese_col = find_header_col(sheet, "Name In Chinese", default=4)
    sender_col = find_header_col(sheet, "Sender", default=2)
    desc_col = find_header_col(sheet, "Description", default=16)
    signals = []
    for row in range(2, sheet.max_row + 1):
        chinese = normalize(sheet.cell(row, chinese_col).value)
        signal = normalize(sheet.cell(row, signal_col).value)
        if not chinese or not signal:
            continue
        signals.append(
            RequirementSignal(
                row=row,
                signal_name=signal,
                chinese_name=chinese,
                sender=normalize(sheet.cell(row, sender_col).value),
                description=normalize(sheet.cell(row, desc_col).value),
            )
        )
    return signals


def compact_key(text):
    return re.sub(r"\s+", "", normalize(text))


def signal_to_candidate(signal, include_description=True):
    candidate = {
        "row": signal.row,
        "signalName": signal.signal_name,
        "chineseName": signal.chinese_name,
        "sender": signal.sender,
    }
    if include_description:
        candidate["description"] = signal.description[:300]
    return candidate


def comparable_name(text):
    return re.sub(r"[\s\u3000，,。；;：:（）()\[\]【】<>《》/\\_-]+", "", normalize(text)).lower()


def strip_value_words(text):
    text = comparable_name(text)
    text = re.sub(r"[+-]?\d+(?:\.\d+)?(?:kmh|km/h|mm|%)?", "", text, flags=re.I)
    text = re.sub(r"[xX]?[0-9a-fA-F]+", "", text)
    for word in [
        "指定",
        "对应",
        "保持为",
        "调整为",
        "设置为",
        "变为",
        "等于",
        "为",
        "仍然",
        "继续",
        "实时计算",
        "调整范围",
        "标准",
        "运动",
        "维修",
        "一键调平",
        "舒适",
        "自适应",
        "p挡",
        "p档",
    ]:
        text = text.replace(word, "")
    return text


def stable_signal_choice_key(role, target_names, search_queries, signal_profile, max_matches):
    names = [strip_value_words(item) for item in target_names + search_queries]
    names = sorted({item for item in names if item})
    if not names:
        return None
    kind = normalize((signal_profile or {}).get("kind")).lower()
    positions = tuple(extract_profile_positions(signal_profile or {}))
    return ("stable_signal_choice", role, max_matches, kind, positions, tuple(names))


def stable_cache_get(cache, key):
    value = cache.get(key)
    if isinstance(value, dict) and value.get("confirmed"):
        names = value.get("names") or ()
        return tuple(names) if isinstance(names, (list, tuple)) else ()
    return None


def stable_cache_set(cache, key, names, confirmed=False):
    if key and names:
        cache[key] = {"confirmed": bool(confirmed), "names": tuple(names)}


def remember_stable_choice(cache, key, names):
    if not key or not names:
        return
    names = tuple(names)
    current = cache.get(key)
    if isinstance(current, dict):
        if current.get("confirmed"):
            return
        if tuple(current.get("names") or ()) == names:
            stable_cache_set(cache, key, names, confirmed=True)
            return
    stable_cache_set(cache, key, names, confirmed=False)


def chunks(items, size):
    for start in range(0, len(items), size):
        yield items[start : start + size]


def set_cell_value(sheet, row, col, value):
    cell = sheet.cell(row, col)
    if isinstance(cell, MergedCell):
        return
    cell.value = value


def chat_completion(messages, model, base_url, api_key, timeout):
    body = json.dumps(
        {
            "model": model,
            "messages": messages,
            "temperature": 0.0,
            "response_format": {"type": "json_object"},
        },
        ensure_ascii=False,
    ).encode("utf-8")
    request = urllib.request.Request(
        base_url.rstrip("/") + "/chat/completions",
        data=body,
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        payload = json.loads(response.read().decode("utf-8"))
    return payload["choices"][0]["message"]["content"]


def extract_json(text):
    start = text.find("{")
    end = text.rfind("}")
    if start >= 0 and end >= start:
        text = text[start : end + 1]
    return json.loads(text)


def llm_extract_target_names(phrase, role, model, base_url, api_key, timeout, cache):
    key = ("target_names", compact_key(phrase), role)
    if key in cache:
        return cache[key]
    messages = [
        {"role": "system", "content": EXTRACT_TARGET_PROMPT},
        {"role": "user", "content": json.dumps({"phrase": phrase, "actionRole": role}, ensure_ascii=False)},
    ]
    content = chat_completion(messages, model, base_url, api_key, timeout)
    data = extract_json(content)
    names = []
    for name in data.get("targetChineseNames", []):
        name = normalize(name)
        if name and name not in names:
            names.append(name)
    search_queries = []
    for query in data.get("searchQueries", []):
        query = normalize(query)
        if query and query not in search_queries:
            search_queries.append(query)
    signal_profile = data.get("signalProfile") or {}
    if not isinstance(signal_profile, dict):
        signal_profile = {}
    result = {
        "targetNames": names,
        "searchQueries": search_queries,
        "signalProfile": signal_profile,
    }
    cache[key] = result
    time.sleep(0.12)
    return result


def find_signals_by_target_names(target_names, signals):
    matches = []
    seen = set()
    indexed = [(signal, comparable_name(signal.chinese_name)) for signal in signals]
    for target_name in target_names:
        target = comparable_name(target_name)
        if not target:
            continue
        exact = [signal for signal, chinese in indexed if chinese == target]
        contains = [
            signal
            for signal, chinese in indexed
            if chinese != target and (target in chinese or chinese in target)
        ]
        for signal in exact + contains:
            if signal.signal_name not in seen:
                seen.add(signal.signal_name)
                matches.append(signal)
    return matches


def classify_signal_kind(signal):
    signal_name = normalize(getattr(signal, "signal_name", ""))
    name_text = normalize(getattr(signal, "chinese_name", ""))
    signal_lower = signal_name.lower()
    text = f"{signal_name} {name_text}".lower()

    if re.search(r"(vld|valid|enable|en)$|(^|_)(vld|valid|enable|en)($|_)", signal_lower):
        return "valid"
    if re.search(r"(fault|err|error|fail|dtc)", signal_lower):
        return "fault"
    if re.search(r"(fb|feedback|actual|act)", signal_lower):
        return "actual_value"
    if re.search(r"(req|request|cmd)$|(^|_)(req|request|cmd)($|_)", signal_lower):
        return "request"

    if re.search(r"有效|使能|可用|就绪", name_text):
        return "valid"
    if re.search(r"故障|错误|异常|失效", name_text):
        return "fault"
    if re.search(r"实际|反馈|当前|检测|实际值", name_text):
        return "actual_value"
    if re.search(r"请求|设定|目标|命令", name_text):
        return "request"
    if re.search(r"\b(mode)\b|模式|状态|档位|工况", text):
        return "status"
    return "unknown"


def extract_profile_positions(profile):
    positions = profile.get("positions") or []
    if isinstance(positions, str):
        positions = [positions]
    normalized = []
    for item in positions:
        item = normalize(item).upper()
        if item and item not in normalized:
            normalized.append(item)
    return normalized


def candidate_matches_profile(signal, profile, phrase):
    if not profile:
        return True
    phrase_text = normalize(phrase)
    signal_text = " ".join(
        [
            normalize(getattr(signal, "signal_name", "")),
            normalize(getattr(signal, "chinese_name", "")),
        ]
    )
    target_kind = normalize(profile.get("kind")).lower()
    candidate_kind = classify_signal_kind(signal)
    if target_kind and candidate_kind != "unknown":
        if target_kind in {"actual_value", "request", "valid", "fault", "status", "target"} and candidate_kind != target_kind:
            return False
    positions = extract_profile_positions(profile)
    if positions:
        pos_text = signal_text.upper()
        if not any(pos in pos_text for pos in positions):
            return False
    object_hint = normalize(profile.get("object"))
    if object_hint:
        compact_object = comparable_name(object_hint)
        compact_signal = comparable_name(signal_text)
        if compact_object and compact_object not in compact_signal and compact_signal not in compact_object:
            phrase_compact = comparable_name(phrase_text)
            if compact_object not in phrase_compact and phrase_compact not in compact_object:
                return False
    return True


def filter_candidates_by_profile(candidates, profile, phrase):
    if not profile:
        return candidates
    filtered = [signal for signal in candidates if candidate_matches_profile(signal, profile, phrase)]
    return filtered or candidates


def candidate_rank_score(signal, phrase, role, target_names, search_queries):
    signal_kind = classify_signal_kind(signal)
    signal_name = normalize(signal.signal_name)
    chinese = comparable_name(signal.chinese_name)
    phrase_key = comparable_name(phrase)
    phrase_text = normalize(phrase)
    queries = [comparable_name(item) for item in target_names + search_queries + [phrase]]
    queries = [item for item in queries if item]

    score = 0
    phrase_focus = []
    for word in ["模式", "高度", "请求", "反馈", "状态", "有效", "故障", "车速", "挡", "电压", "电流", "压力", "温度"]:
        if word in phrase_text:
            phrase_focus.append(word)
    signal_focus = normalize(signal.chinese_name)
    signal_name_focus = signal_name.lower()
    for query in queries:
        if query == chinese:
            score += 120
        elif query and chinese and (query in chinese or chinese in query):
            score += 70
        if query and chinese:
            query_bigrams = {query[index : index + 2] for index in range(len(query) - 1)}
            chinese_bigrams = {chinese[index : index + 2] for index in range(len(chinese) - 1)}
            score += len(query_bigrams & chinese_bigrams) * 6
    if phrase_key and chinese and chinese in phrase_key:
        score += 50
    for word in phrase_focus:
        if word in signal_focus:
            score += 40
        else:
            score -= 15
    if any(word in phrase_text for word in ["模式", "状态"]):
        if "模式" in signal_focus:
            score += 25
        if any(word in signal_name_focus for word in ["ctrl", "req"]) and "fb" not in signal_name_focus:
            score += 10
        if any(word in signal_name_focus for word in ["valve", "dir", "sw", "ctrl"]) and "fb" not in signal_name_focus and "sts" not in signal_name_focus:
            score -= 20
    if any(word in phrase_text for word in ["请求", "设置", "输入"]):
        if any(word in signal_name_focus for word in ["req", "request", "cmd"]):
            score += 30
        if any(word in signal_name_focus for word in ["ctrl", "act", "fb", "sts"]) and "req" not in signal_name_focus:
            score -= 25
    if re.search(r"[PRND]挡|[PRND]档|挡位|档位|换挡|gear", phrase_text, re.I):
        if re.search(r"挡|档|换挡", signal.chinese_name) or re.search(r"gear|lever|shift", signal_name_focus):
            score += 160
        else:
            score -= 60
    if re.search(r"禁用|禁止|关闭|停用", signal.chinese_name) and not re.search(r"禁用|禁止|关闭|停用", phrase):
        score -= 45

    signal_lower = signal_name.lower()
    if role == "write":
        if signal_kind == "request":
            score += 90
        if signal_kind in {"actual_value", "valid", "fault", "status"}:
            score -= 90
        if re.search(r"(req|request|cmd)$", signal_lower):
            score += 80
        if re.search(r"(fb|feedback|actual|act|vld|valid|fault|err|error)", signal_lower):
            score -= 120
    elif role == "read":
        if signal_kind in {"actual_value", "status"}:
            score += 50
        if signal_kind == "request":
            score -= 80
        if re.search(r"(fb|feedback|actual|act)$", signal_lower):
            score += 60
        if re.search(r"(req|request|cmd)$", signal_lower):
            score -= 100

    # Prefer a specific matched concept over a broad control/status signal when
    # both are present in the candidate pool.
    if len(chinese) <= 4:
        score -= 10
    return score


def char_ngrams(text, min_n=2, max_n=3):
    text = comparable_name(text)
    grams = set()
    for n in range(min_n, max_n + 1):
        grams.update(text[index : index + n] for index in range(max(0, len(text) - n + 1)))
    return {gram for gram in grams if gram}


def lexical_overlap_score(left, right):
    left_grams = char_ngrams(left)
    right_grams = char_ngrams(right)
    if not left_grams or not right_grams:
        return 0.0
    return len(left_grams & right_grams) / len(left_grams)


def prune_weak_name_candidates(candidates, target_names, search_queries):
    anchors = [item for item in target_names + search_queries if comparable_name(item)]
    if not anchors or len(candidates) <= 1:
        return candidates
    scored = []
    for signal in candidates:
        best = max(lexical_overlap_score(anchor, signal.chinese_name) for anchor in anchors)
        scored.append((best, signal))
    best_score = max(score for score, _ in scored)
    if best_score <= 0:
        return candidates
    threshold = max(0.18, best_score * 0.45)
    filtered = [signal for score, signal in scored if score >= threshold]
    return filtered or candidates


def order_candidates(candidates, phrase, role, target_names, search_queries, limit=None):
    seen = set()
    unique = []
    for signal in candidates:
        if signal.signal_name in seen:
            continue
        seen.add(signal.signal_name)
        unique.append(signal)
    if not re.search(r"[PRND]挡|[PRND]档|挡位|档位|换挡|gear", normalize(phrase), re.I):
        unique = prune_weak_name_candidates(unique, target_names, search_queries)
    unique.sort(
        key=lambda signal: candidate_rank_score(signal, phrase, role, target_names, search_queries),
        reverse=True,
    )
    return unique[:limit] if limit else unique


POSITION_CODES = ("FL", "FR", "RL", "RR")


def phrase_has_specific_position(phrase):
    text = normalize(phrase).upper()
    if any(pos in text for pos in POSITION_CODES):
        return True
    return bool(re.search(r"左前|右前|左后|右后|前左|前右|后左|后右", normalize(phrase)))


def positional_signature(signal_name):
    for pos in POSITION_CODES:
        if pos in signal_name:
            return signal_name.replace(pos, "{POS}", 1), pos
    return None, None


def expand_positional_matches(matches, signals, phrase, max_matches):
    if phrase_has_specific_position(phrase) or max_matches <= 1:
        return matches
    signal_by_name = {signal.signal_name: signal for signal in signals}
    expanded = []
    seen = set()
    for match in matches:
        signal_name = match["signalName"]
        signature, _ = positional_signature(signal_name)
        siblings = []
        if signature:
            siblings = [
                signature.replace("{POS}", pos)
                for pos in POSITION_CODES
                if signature.replace("{POS}", pos) in signal_by_name
            ]
        if len(siblings) >= 2:
            for sibling in siblings:
                if sibling not in seen:
                    seen.add(sibling)
                    expanded.append({"signalName": sibling, "value": match["value"]})
        elif signal_name not in seen:
            seen.add(signal_name)
            expanded.append(match)
    return expanded[:max_matches]


def llm_select_from_candidates(
    phrase,
    role,
    candidates,
    model,
    base_url,
    api_key,
    timeout,
    max_matches,
    case_context="",
    include_description=False,
):
    payload = {
        "phrase": phrase,
        "caseContext": case_context,
        "actionRole": role,
        "maxMatches": max_matches,
        "selectionInstruction": "??? candidates.signalName ? candidates.chineseName ????????? Description ??????",
        "candidates": [signal_to_candidate(item, include_description=include_description) for item in candidates],
    }
    if include_description:
        payload["valueInstruction"] = "??????????????? description ?? value????? description ??????????? write#/read# ????????? description ?? 0x/??????????? 0x/???????????"
    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
    ]
    content = chat_completion(messages, model, base_url, api_key, timeout)
    data = extract_json(content)
    matches = []
    known_signal_names = {item.signal_name for item in candidates}
    for item in data.get("matches", []):
        signal_name = normalize(item.get("signalName"))
        value = normalize(item.get("value"))
        if signal_name in known_signal_names and signal_name not in {m["signalName"] for m in matches}:
            matches.append({"signalName": signal_name, "value": value})
        if len(matches) >= max_matches:
            break
    time.sleep(0.12)
    return matches


def llm_infer_value(phrase, role, signal, model, base_url, api_key, timeout, case_context=""):
    if not signal:
        return ""
    matches = llm_select_from_candidates(
        phrase,
        role,
        [signal],
        model,
        base_url,
        api_key,
        timeout,
        1,
        case_context,
        include_description=True,
    )
    if matches:
        return matches[0].get("value", "")
    return ""


def resolve_match_values(matches, signals, phrase, role, model, base_url, api_key, timeout, case_context=""):
    signal_by_name = {signal.signal_name: signal for signal in signals}
    resolved = []
    for match in matches:
        signal_name = match["signalName"]
        signal = signal_by_name.get(signal_name)
        value = llm_infer_value(phrase, role, signal, model, base_url, api_key, timeout, case_context)
        resolved.append({"signalName": signal_name, "value": value})
    return resolved


def semantic_match_key(role, target_names, search_queries, signal_profile, max_matches):
    profile_kind = normalize((signal_profile or {}).get("kind")).lower()
    positions = extract_profile_positions(signal_profile or {})
    names = sorted({comparable_name(item) for item in target_names + search_queries if comparable_name(item)})
    if not names:
        return None
    return (
        "signal_choice",
        role,
        max_matches,
        profile_kind,
        tuple(positions),
        tuple(names),
    )


def matches_from_signal_names(signal_names):
    return [{"signalName": signal_name, "value": ""} for signal_name in signal_names]


def signal_names_from_matches(matches):
    names = []
    for match in matches:
        signal_name = match["signalName"]
        if signal_name and signal_name not in names:
            names.append(signal_name)
    return tuple(names)


def should_read_multiple_signals(phrase):
    text = normalize(phrase).upper()
    if any(pos in text for pos in POSITION_CODES):
        return True
    if re.search(r"四轮|四个|多个|各轮|前后|左右|左前|右前|左后|右后|FL|FR|RL|RR", text):
        return True
    return "高度" in normalize(phrase)


def max_matches_for_phrase(phrase, role):
    if role == "write":
        return 1
    return 4 if should_read_multiple_signals(phrase) else 1


def role_adjusted_profile(profile, role):
    adjusted = dict(profile or {})
    if role == "write":
        adjusted["kind"] = "request"
    return adjusted


def llm_match_signal(
    phrase,
    role,
    signals,
    model,
    base_url,
    api_key,
    timeout,
    cache,
    max_matches=4,
    batch_size=60,
    retriever=None,
    embedding_top_k=20,
    case_context="",
):
    key = (
        compact_key(phrase),
        compact_key(case_context),
        role,
        max_matches,
        batch_size,
        bool(retriever),
        embedding_top_k,
    )
    if key in cache:
        return cache[key]

    extraction = llm_extract_target_names(phrase, role, model, base_url, api_key, timeout, cache)
    target_names = extraction["targetNames"]
    search_queries = extraction["searchQueries"]
    if not target_names and not search_queries:
        result = {
            "matches": [],
            "targetNames": target_names,
            "searchQueries": search_queries,
            "signalProfile": extraction.get("signalProfile") or {},
            "strongCandidateCount": 0,
            "embeddingCandidateCount": 0,
        }
        cache[key] = result
        return result
    signal_profile = role_adjusted_profile(extraction.get("signalProfile") or {}, role)
    choice_key = semantic_match_key(role, target_names, search_queries, signal_profile, max_matches)
    stable_choice_key = stable_signal_choice_key(role, target_names, search_queries, signal_profile, max_matches)
    cached_signal_names = stable_cache_get(cache, stable_choice_key) if stable_choice_key else None
    if cached_signal_names:
        final_matches = resolve_match_values(
            matches_from_signal_names(cached_signal_names),
            signals,
            phrase,
            role,
            model,
            base_url,
            api_key,
            timeout,
            case_context,
        )
        result = {
            "matches": final_matches,
            "targetNames": target_names,
            "searchQueries": search_queries,
            "signalProfile": signal_profile,
            "strongCandidateCount": 0,
            "embeddingCandidateCount": 0,
        }
        cache[key] = result
        return result

    strong_candidates = filter_candidates_by_profile(
        find_signals_by_target_names(target_names, signals),
        signal_profile,
        phrase,
    )

    embedding_candidates = []
    if retriever:
        query_parts = target_names + search_queries + [phrase]
        query = " ".join(part for part in query_parts if normalize(part))
        embedding_hits = retriever.search(query, top_k=embedding_top_k)
        embedding_candidates = filter_candidates_by_profile(
            [item["signal"] for item in embedding_hits],
            signal_profile,
            phrase,
        )

    merged_candidates = order_candidates(
        strong_candidates + embedding_candidates,
        phrase,
        role,
        target_names,
        search_queries,
        limit=max(embedding_top_k, 12),
    )
    if merged_candidates:
        if max_matches == 1:
            final_matches = [{"signalName": merged_candidates[0].signal_name, "value": ""}]
        else:
            final_matches = llm_select_from_candidates(
                phrase,
                role,
                merged_candidates,
                model,
                base_url,
                api_key,
                timeout,
                max_matches,
                case_context,
            )
        final_matches = resolve_match_values(
            final_matches,
            signals,
            phrase,
            role,
            model,
            base_url,
            api_key,
            timeout,
            case_context,
        )
        final_matches = expand_positional_matches(final_matches, signals, phrase, max_matches)
        if final_matches:
            names = signal_names_from_matches(final_matches)
            if choice_key:
                cache[choice_key] = names
            if stable_choice_key:
                remember_stable_choice(cache, stable_choice_key, names)
        result = {
            "matches": final_matches,
            "targetNames": target_names,
            "searchQueries": search_queries,
            "signalProfile": signal_profile,
            "strongCandidateCount": len(strong_candidates),
            "embeddingCandidateCount": len(embedding_candidates),
        }
        cache[key] = result
        return result

    batch_matches = []
    for batch in chunks(signals, batch_size):
        batch_matches.extend(
            llm_select_from_candidates(
                phrase,
                role,
                order_candidates(
                    filter_candidates_by_profile(batch, signal_profile, phrase),
                    phrase,
                    role,
                    target_names,
                    search_queries,
                ),
                model,
                base_url,
                api_key,
                timeout,
                max_matches,
                case_context,
            )
        )

    deduped = []
    seen = set()
    signal_by_name = {signal.signal_name: signal for signal in signals}
    for match in batch_matches:
        signal_name = match["signalName"]
        if signal_name in signal_by_name and signal_name not in seen:
            seen.add(signal_name)
            deduped.append(match)

    if len(deduped) > max_matches:
        finalist_signals = [signal_by_name[item["signalName"]] for item in deduped]
        if max_matches == 1:
            ranked_finalists = order_candidates(finalist_signals, phrase, role, target_names, search_queries)
            final_matches = [{"signalName": ranked_finalists[0].signal_name, "value": ""}] if ranked_finalists else []
        else:
            final_matches = llm_select_from_candidates(
                phrase,
                role,
                finalist_signals,
                model,
                base_url,
                api_key,
                timeout,
                max_matches,
                case_context,
            )
        final_matches = resolve_match_values(
            final_matches,
            signals,
            phrase,
            role,
            model,
            base_url,
            api_key,
            timeout,
            case_context,
        )
        final_matches = expand_positional_matches(final_matches, signals, phrase, max_matches)
    else:
        final_matches = resolve_match_values(
            deduped,
            signals,
            phrase,
            role,
            model,
            base_url,
            api_key,
            timeout,
            case_context,
        )
        final_matches = expand_positional_matches(final_matches, signals, phrase, max_matches)

    if final_matches:
        names = signal_names_from_matches(final_matches)
        if choice_key:
            cache[choice_key] = names
        if stable_choice_key:
            remember_stable_choice(cache, stable_choice_key, names)
    result = {
        "matches": final_matches,
        "targetNames": target_names,
        "searchQueries": search_queries,
        "signalProfile": signal_profile,
        "strongCandidateCount": 0,
        "embeddingCandidateCount": len(embedding_candidates),
    }
    cache[key] = result
    return result

def fixed_command(item):
    compact = compact_key(item)
    if "初始状态检测" in compact or "整车上电" in compact or re.search(r"power\s*up|powerup", compact, re.I):
        return ["POWERUP"]
    if "整车下电" in compact or re.search(r"power\s*down|powerdown", compact, re.I):
        return ["POWERDOWN"]
    return []


def signal_role_for_item(role, item):
    if role == "write" and re.search(r"^(检测|检查|确认|读取|观测|观察|判断)", normalize(item)):
        return "read"
    return role


def render_commands(
    text,
    role,
    case_details,
    signals,
    llm_args,
    cache,
    audit_rows,
    sheet_name,
    row,
    source_col,
    retriever=None,
    embedding_top_k=20,
):
    commands = []
    for item in split_items(text):
        fixed = fixed_command(item)
        if fixed:
            commands.extend(fixed)
            continue
        item_role = signal_role_for_item(role, item)
        if item_role == "initial":
            continue

        match_result = llm_match_signal(
            item,
            item_role,
            signals,
            *llm_args,
            cache,
            max_matches=max_matches_for_phrase(item, item_role),
            retriever=retriever,
            embedding_top_k=embedding_top_k,
            case_context=normalize(case_details),
        )
        matched_signals = match_result["matches"]
        target_names = " | ".join(match_result["targetNames"])
        search_queries = " | ".join(match_result["searchQueries"])
        strong_count = match_result["strongCandidateCount"]
        embedding_count = match_result["embeddingCandidateCount"]
        if not matched_signals:
            audit_rows.append([
                sheet_name,
                row,
                source_col,
                item_role,
                item,
                target_names,
                search_queries,
                strong_count,
                embedding_count,
                "",
                "",
                "unmatched",
            ])
            continue

        if item_role == "write":
            for matched in matched_signals:
                value = matched["value"] or "指定值"
                commands.append(f"write#{matched['signalName']}={value}")
                audit_rows.append([
                    sheet_name,
                    row,
                    source_col,
                    role,
                    item,
                    target_names,
                    search_queries,
                    strong_count,
                    embedding_count,
                    matched["signalName"],
                    value,
                    "write",
                ])
        else:
            for matched in matched_signals:
                value = matched["value"]
                if value:
                    commands.append(f"read#{matched['signalName']}={value}")
                else:
                    commands.append(f"read#{matched['signalName']}={item}")
                audit_rows.append([
                    sheet_name,
                    row,
                    source_col,
                    role,
                    item,
                    target_names,
                    search_queries,
                    strong_count,
                    embedding_count,
                    matched["signalName"],
                    value,
                    "read",
                ])
    return commands


def fill_workbook(
    case_path,
    requirement_path,
    output_path,
    model,
    base_url,
    api_key,
    timeout,
    max_sheets=None,
    embedding_model=None,
    embedding_cache=None,
    embedding_top_k=20,
    embedding_dimensions=1024,
    use_embedding=True,
):
    signals = load_requirement_signals(requirement_path)
    if not signals:
        raise ValueError(f"No usable SignalName values found in {requirement_path}")
    retriever = None
    if use_embedding:
        retriever = EmbeddingRetriever.build(
            signals,
            embedding_model,
            base_url,
            api_key,
            timeout,
            embedding_cache,
            embedding_dimensions,
        )
    workbook = load_workbook(case_path)
    cache = {}
    audit_rows = []
    llm_args = (model, base_url, api_key, timeout)

    for sheet in workbook.worksheets:
        if sheet.title == "Qwen匹配审计":
            continue
        hil_cols = find_hil_columns(sheet)
        for col in hil_cols.values():
            for row in range(2, sheet.max_row + 1):
                set_cell_value(sheet, row, col, None)

    sheets = workbook.worksheets[:max_sheets] if max_sheets else workbook.worksheets
    for sheet in sheets:
        hil_cols = find_hil_columns(sheet)
        text_cols = text_columns_from_hil(hil_cols)
        for start, end in case_ranges(sheet):
            initial_text = sheet.cell(start, text_cols["initial"]).value
            initial_commands = render_commands(
                initial_text,
                "initial",
                "",
                signals,
                llm_args,
                cache,
                audit_rows,
                sheet.title,
                start,
                "initial",
                retriever,
                embedding_top_k,
            )
            set_cell_value(sheet, start, hil_cols["initial"], "\n".join(initial_commands))

            for row in range(start, end + 1):
                step = sheet.cell(row, text_cols["step"]).value
                if step:
                    set_cell_value(sheet, row, hil_cols["step"], step)

                case_details = sheet.cell(row, 4).value or sheet.cell(start, 4).value
                action_commands = render_commands(
                    sheet.cell(row, text_cols["action"]).value,
                    "write",
                    case_details,
                    signals,
                    llm_args,
                    cache,
                    audit_rows,
                    sheet.title,
                    row,
                    "action",
                    retriever,
                    embedding_top_k,
                )
                expected_commands = render_commands(
                    sheet.cell(row, text_cols["expected"]).value,
                    "read",
                    case_details,
                    signals,
                    llm_args,
                    cache,
                    audit_rows,
                    sheet.title,
                    row,
                    "expected",
                    retriever,
                    embedding_top_k,
                )
                if action_commands:
                    set_cell_value(sheet, row, hil_cols["action"], "\n".join(action_commands))
                else:
                    set_cell_value(sheet, row, hil_cols["action"], None)
                if expected_commands:
                    set_cell_value(sheet, row, hil_cols["expected"], "\n".join(expected_commands))
                else:
                    set_cell_value(sheet, row, hil_cols["expected"], None)

        for col in hil_cols.values():
            for row in range(1, sheet.max_row + 1):
                sheet.cell(row, col).font = Font(name="Arial", size=10)

    if "Qwen匹配审计" in workbook.sheetnames:
        del workbook["Qwen匹配审计"]
    audit_sheet = workbook.create_sheet("Qwen匹配审计")
    audit_sheet.append([
        "sheet",
        "row",
        "sourceColumn",
        "role",
        "sourceText",
        "targetChineseNames",
        "searchQueries",
        "strongCandidateCount",
        "embeddingCandidateCount",
        "SignalName",
        "value",
        "status",
    ])
    for audit_row in audit_rows:
        audit_sheet.append(audit_row)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path, len(cache)


def main():
    parser = argparse.ArgumentParser(description="Use an LLM to fill HIL signal columns from text cases.")
    parser.add_argument("--case", default=str(default_file("用例1-1.xlsx")))
    parser.add_argument("--requirement", default=str(default_requirement_file()))
    parser.add_argument("--output", default=str(PROJECT_ROOT / "demo" / "用例1-1_HIL_Qwen.xlsx"))
    parser.add_argument("--provider", default=DEFAULT_PROVIDER, choices=sorted(PROVIDER_DEFAULTS.keys()))
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL)
    parser.add_argument("--api-key", default=DEFAULTS["api_key"] or os.environ.get("OPENAI_API_KEY") or os.environ.get("QWEN_API_KEY") or os.environ.get("DASHSCOPE_API_KEY"))
    parser.add_argument("--timeout", type=int, default=90)
    parser.add_argument("--max-sheets", type=int, default=None)
    parser.add_argument("--embedding-model", default=DEFAULT_EMBEDDING_MODEL)
    parser.add_argument("--embedding-top-k", type=int, default=20)
    parser.add_argument("--embedding-dimensions", type=int, default=1024)
    parser.add_argument("--embedding-cache", default=str(PROJECT_ROOT / "demo" / ".qwen_embedding_cache.json"))
    parser.add_argument("--no-embedding", action="store_true")
    args = parser.parse_args()

    if not args.api_key:
        raise SystemExit("Please set the API key for the selected provider, or pass --api-key.")

    output, matched_phrases = fill_workbook(
        Path(args.case),
        Path(args.requirement),
        Path(args.output),
        args.model,
        args.base_url,
        args.api_key,
        args.timeout,
        args.max_sheets,
        args.embedding_model,
        Path(args.embedding_cache),
        args.embedding_top_k,
        args.embedding_dimensions,
        not args.no_embedding,
    )
    print(f"Matched phrases: {matched_phrases}")
    print(f"Output: {output.resolve()}")


if __name__ == "__main__":
    main()
