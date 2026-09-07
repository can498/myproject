"""Repair obvious HIL formatting/signal issues in a generated case workbook.

This is a post-generation cleanup for already-created xlsx files. It does not
call an LLM and does not replace the main converter.
"""

import argparse
import re
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

HEIGHT_SIGNAL_REPLACEMENTS = {
    "SuspensionHeightFL": "ECASActLvlFL",
    "SuspensionHeightFR": "ECASActLvlFR",
    "SuspensionHeightRL": "ECASActLvlRL",
    "SuspensionHeightRR": "ECASActLvlRR",
}


def normalize_text(value):
    if value is None:
        return ""
    return str(value).replace("\u3000", " ").strip()


def find_hil_columns(sheet):
    for col in range(1, sheet.max_column + 1):
        value = sheet.cell(1, col).value
        if isinstance(value, str) and "HIL" in value:
            return {
                "initial": col,
                "step_id": col + 1,
                "action": col + 2,
                "expected": col + 3,
            }
    return {"initial": 10, "step_id": 11, "action": 12, "expected": 13}


def find_text_columns(sheet):
    hil = find_hil_columns(sheet)
    return {
        "action": max(1, hil["action"] - 4),
        "expected": max(1, hil["expected"] - 4),
    }


def normalize_power_command(cell_value):
    text = normalize_text(cell_value)
    if not text:
        return cell_value
    lines = []
    changed = False
    for line in re.split(r"(\r?\n)", text):
        compact = re.sub(r"[\s_;_-]+", "", line).upper()
        if compact in ("POWERDOWN", "POWEROFF"):
            lines.append("POWER DOWN")
            changed = True
        else:
            lines.append(line)
    return "".join(lines) if changed else cell_value


def replace_height_signals(cell_value):
    text = normalize_text(cell_value)
    if not text:
        return cell_value
    changed = False
    for old, new in HEIGHT_SIGNAL_REPLACEMENTS.items():
        if old in text:
            text = text.replace(old, new)
            changed = True
    return text if changed else cell_value


def source_gear_letter(action_text, expected_text):
    source = normalize_text(action_text) + "\n" + normalize_text(expected_text)
    for letter in ("P", "R", "N", "D"):
        if f"\u6302{letter}\u6321" in source or f"\u6302{letter}\u6863" in source or f"{letter}\u6321" in source or f"{letter}\u6863" in source:
            return letter
    return ""


def repair_gear_commands(cell_value, gear_letter):
    if not gear_letter:
        return cell_value
    text = normalize_text(cell_value)
    if not text:
        return cell_value
    original = text
    text = re.sub(r"(write#GearLever\s*=\s*)[0-9]+", rf"\1{gear_letter}", text)
    text = re.sub(r"(read#GearLever\s*==\s*)[0-9]+", rf"\1{gear_letter}", text)
    text = re.sub(r"read#ModeSts\s*==\s*[0-9A-Za-z]+", f"read#GearLever=={gear_letter}", text)
    return text if text != original else cell_value


def should_expand_actual_height(expected_text):
    text = normalize_text(expected_text)
    return "\u5b9e\u9645\u7a7a\u7c27\u9ad8\u5ea6" in text


def expand_single_height_read(cell_value, expected_text):
    if not should_expand_actual_height(expected_text):
        return cell_value
    text = normalize_text(cell_value)
    if not text:
        return cell_value
    if any(signal in text for signal in ("ECASActLvlFR", "ECASActLvlRL", "ECASActLvlRR")):
        return cell_value

    pattern = re.compile(r"read#ECASActLvlFL\s*(==|=)\s*([^;\n\r]+)")
    match = pattern.search(text)
    if not match:
        return cell_value

    op = match.group(1)
    value = match.group(2).strip()
    expanded = "\n".join([
        f"read#ECASActLvlFL{op}{value}",
        f"read#ECASActLvlFR{op}{value}",
        f"read#ECASActLvlRL{op}{value}",
        f"read#ECASActLvlRR{op}{value}",
    ])
    return pattern.sub(expanded, text, count=1)


def repair_workbook(input_path, output_path):
    workbook = load_workbook(input_path)
    changes = []
    for sheet in workbook.worksheets:
        hil = find_hil_columns(sheet)
        text_cols = find_text_columns(sheet)
        for row in range(1, sheet.max_row + 1):
            action_text = sheet.cell(row, text_cols["action"]).value
            expected_text = sheet.cell(row, text_cols["expected"]).value
            gear = source_gear_letter(action_text, expected_text)
            for col_name in ("initial", "action", "expected"):
                col = hil[col_name]
                old_value = sheet.cell(row, col).value
                new_value = normalize_power_command(old_value)
                new_value = replace_height_signals(new_value)
                new_value = repair_gear_commands(new_value, gear)
                if col_name == "expected":
                    new_value = expand_single_height_read(new_value, expected_text)
                if new_value != old_value:
                    sheet.cell(row, col).value = new_value
                    changes.append((sheet.title, row, col_name, old_value, new_value))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return changes


def default_input():
    candidates = [
        p for p in (PROJECT_ROOT / "demo").glob("*HIL_AI_sheet1.xlsx")
        if not p.name.startswith("~$")
    ]
    if candidates:
        return str(candidates[0])
    raise FileNotFoundError("Cannot find generated HIL sheet1 workbook.")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", default=default_input())
    parser.add_argument("--output", default=str(PROJECT_ROOT / "demo" / "用例1-1_HIL_AI_sheet1_修复.xlsx"))
    args = parser.parse_args()

    changes = repair_workbook(Path(args.input), Path(args.output))
    print(f"Repaired cells: {len(changes)}")
    print(f"Output: {Path(args.output).resolve()}")


if __name__ == "__main__":
    main()
