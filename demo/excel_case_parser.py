"""批量解析《用例.xlsx》中所有工作表的HIL测试用例。"""

import argparse
import copy
import json
import re
from pathlib import Path

from openpyxl import load_workbook


WRITE_RE = re.compile(r"^write#\s*([^=]+?)\s*=\s*(.+?)\s*;?$", re.I)
WAIT_RE = re.compile(r"^wait\s*=\s*(\d+(?:\.\d+)?)\s*(ms|s|min|h)\s*;?$", re.I)
READ_RE = re.compile(
    r"^read#\s*([A-Za-z_][A-Za-z0-9_]*)\s*(==|!=|>=|<=|>|<|=)\s*(.+?)\s*;?$",
    re.I,
)
READ_ONLY_RE = re.compile(r"^read#\s*([A-Za-z_][A-Za-z0-9_]*)(.*)$", re.I)

OPERATOR_MAP = {
    "=": "Equal",
    "==": "Equal",
    "!=": "NotEqual",
    ">": "GreaterThan",
    ">=": "GreaterThanOrEqual",
    "<": "LessThan",
    "<=": "LessThanOrEqual",
}

SIGNAL_ALIASES = {"TargertVehSpd": "TargetVehSpd"}
STRING_LITERALS = {"P", "R", "N", "D"}
BUILTIN_ACTIONS = {"POWERUP", "POWERDOWN"}


def sanitize_filename(value, limit=70):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value)).strip(" .")
    value = re.sub(r"\s+", "_", value)
    return (value or "unnamed")[:limit]


def normalize_signal(signal, warnings, cell):
    signal = signal.strip()
    normalized = SIGNAL_ALIASES.get(signal, signal)
    if normalized != signal:
        warnings.append(f"{cell}: 信号名 {signal!r} 已规范为 {normalized!r}")
    return normalized


def split_commands(text):
    if text is None or str(text).strip() in {"", "/"}:
        return []
    normalized = str(text).replace("；", ";").strip()
    return [item.strip() for item in re.split(r"[;\r\n]+", normalized) if item.strip()]


def is_number(value):
    try:
        float(str(value))
        return True
    except ValueError:
        return False


def unquote(value):
    value = str(value).strip()
    if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
        return value[1:-1], True
    return value, False


def classify_value(value):
    value, quoted = unquote(value)
    if quoted or value in STRING_LITERALS:
        return value, "String"
    return value, "Number"


def parse_command(command, warnings, cell):
    upper = command.strip().upper()
    if upper in BUILTIN_ACTIONS:
        return [{"action": upper}]

    wait_match = WAIT_RE.match(command)
    if wait_match:
        value = float(wait_match.group(1))
        return [{
            "action": "Wait",
            "time": int(value) if value.is_integer() else value,
            "unit": wait_match.group(2).lower(),
        }]

    write_match = WRITE_RE.match(command)
    if write_match:
        signal = normalize_signal(write_match.group(1), warnings, cell)
        value, value_type = classify_value(write_match.group(2))
        return [{
            "action": "Write",
            "signal": signal,
            "value": value,
            "valueType": value_type,
            "accessMode": "TextValue" if value_type == "String" else "PhysicalValue",
        }]

    read_match = READ_RE.match(command)
    if read_match:
        signal = normalize_signal(read_match.group(1), warnings, cell)
        raw_operator = read_match.group(2)
        expected, expected_type = classify_value(read_match.group(3))
        if raw_operator == "=":
            warnings.append(f"{cell}: Read断言使用单个'='，已按'=='处理")
        return [{
            "action": "Read",
            "signal": signal,
            "expectType": expected_type,
            "operator": OPERATOR_MAP[raw_operator],
            "expectValue": expected,
            "accessMode": "TextValue" if expected_type == "String" else "PhysicalValue",
            "timeout": "0",
            "timeoutUnit": "ms",
        }]

    read_only_match = READ_ONLY_RE.match(command)
    if read_only_match:
        signal = normalize_signal(read_only_match.group(1), warnings, cell)
        note = read_only_match.group(2).strip(" ;：:")
        warnings.append(f"{cell}: Read缺少标准比较表达式，已生成无断言读取：{command!r}")
        return [{
            "action": "Read",
            "signal": signal,
            "expectType": None,
            "accessMode": "PhysicalValue",
            "offlineValue": 0,
            "description": note or "原用例仅要求读取，未定义断言",
        }]

    warnings.append(f"{cell}: 未知动作已生成占位节点：{command!r}")
    return [{
        "action": "Placeholder",
        "name": command,
        "description": "原Excel中的动作尚未映射到确定的TAE组件",
    }]


def parse_cell(text, warnings, cell):
    actions = []
    for command in split_commands(text):
        actions.extend(parse_command(command, warnings, cell))
    return actions


def case_ranges(sheet):
    """A列非空单元格是用例起点；优先使用A列合并区域确定终点。"""
    starts = [row for row in range(3, sheet.max_row + 1) if sheet.cell(row, 1).value]
    merged_by_start = {
        merged.min_row: merged.max_row
        for merged in sheet.merged_cells.ranges
        if merged.min_col == 1 and merged.max_col == 1
    }
    ranges = []
    for index, start in enumerate(starts):
        next_start = starts[index + 1] if index + 1 < len(starts) else sheet.max_row + 1
        end = merged_by_start.get(start, next_start - 1)
        end = min(end, next_start - 1)
        ranges.append((start, end))
    return ranges


def make_placeholder_variables(actions, warnings):
    """为表达式右侧尚未展开的标识符创建0值占位参数，保证离线框图可运行。"""
    names = set()
    mapped_signals = {
        action.get("signal") for action in actions if action.get("signal")
    }
    for action in actions:
        if action.get("action") not in {"Write", "Read"}:
            continue
        key = "value" if action["action"] == "Write" else "expectValue"
        value = action.get(key)
        if not value or action.get("valueType") == "String" or action.get("expectType") == "String":
            continue
        if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", str(value)):
            names.add(str(value))

    variables = []
    for name in sorted(names):
        description = "参数占位值0，待具体场景展开或项目参数定义"
        if name in mapped_signals:
            description = "比较目标占位值0，原用例引用了另一个变量"
        variables.append({
            "name": name,
            "type": "NumberVariableValue",
            "value": 0,
            "description": description,
        })
        warnings.append(f"参数{name!r}未给出具体值，当前离线框图使用占位值0")
    return variables


def find_hil_columns(sheet):
    """Return the HIL initial/substep/action/expected columns for this sheet."""
    for col in range(1, sheet.max_column + 1):
        value = sheet.cell(1, col).value
        if isinstance(value, str) and "HIL" in value:
            return {
                "initial": col,
                "step_id": col + 1,
                "action": col + 2,
                "expected": col + 3,
            }
    return {
        "initial": 8,
        "step_id": 9,
        "action": 10,
        "expected": 11,
    }


def find_test_level_column(sheet):
    for col in range(1, sheet.max_column + 1):
        value = sheet.cell(1, col).value
        if isinstance(value, str) and "Test Level" in value:
            return col
    return 3


def extract_case(sheet, workbook_path, start_row, end_row, case_index):
    warnings = []
    initialization = []
    steps = []
    cleanup = []
    hil_columns = find_hil_columns(sheet)
    test_level_column = find_test_level_column(sheet)

    for command in split_commands(sheet.cell(start_row, hil_columns["initial"]).value):
        parsed = parse_command(command, warnings, f"{sheet.cell(start_row, hil_columns['initial']).coordinate}")
        initialization.extend(parsed)

    for row in range(start_row, end_row + 1):
        step_id = str(sheet.cell(row, hil_columns["step_id"]).value or f"Row{row}")
        operation_cell = sheet.cell(row, hil_columns["action"])
        expected_cell = sheet.cell(row, hil_columns["expected"])
        operation_actions = parse_cell(operation_cell.value, warnings, operation_cell.coordinate)
        expected_actions = parse_cell(expected_cell.value, warnings, expected_cell.coordinate)
        for action in operation_actions + expected_actions:
            action["sourceStep"] = step_id
            if action["action"] == "POWERDOWN":
                cleanup.append(action)
            elif action["action"] == "POWERUP":
                initialization.append(action)
            else:
                steps.append(action)

    all_actions = initialization + steps + cleanup
    variables = make_placeholder_variables(all_actions, warnings)
    required_mappings = sorted({
        action["signal"]
        for action in steps
        if action.get("action") in {"Write", "Read"} and action.get("signal")
    })

    case_id = f"{case_index:03d}_{sanitize_filename(sheet.title, 24)}_R{start_row}"
    return {
        "caseId": case_id,
        "source": {
            "workbook": str(Path(workbook_path).resolve()),
            "sheet": sheet.title,
            "range": f"A{start_row}:{sheet.cell(end_row, hil_columns['expected']).coordinate}",
        },
        "caseName": str(sheet.cell(start_row, 1).value).strip(),
        "testLevel": str(sheet.cell(start_row, test_level_column).value or "").strip(),
        "concreteSettings": str(sheet.cell(start_row, 2).value or "").strip(),
        "offlineMode": True,
        "variables": variables,
        "initialization": initialization,
        "steps": steps,
        "cleanup": cleanup,
        "requiredMappings": required_mappings,
        "warnings": warnings,
    }


def extract_workbook(workbook_path, selected_sheets=None):
    workbook = load_workbook(workbook_path, data_only=True)
    selected = set(selected_sheets or workbook.sheetnames)
    cases = []
    case_index = 1
    for sheet in workbook.worksheets:
        if sheet.title not in selected:
            continue
        for start_row, end_row in case_ranges(sheet):
            cases.append(extract_case(sheet, workbook_path, start_row, end_row, case_index))
            case_index += 1
    return cases


def write_cases(cases, output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest_cases = []
    for case in cases:
        filename = f"{case['caseId']}_{sanitize_filename(case['caseName'])}.json"
        path = output_dir / filename
        path.write_text(json.dumps(case, ensure_ascii=False, indent=2), encoding="utf-8")
        manifest_cases.append({
            "caseId": case["caseId"],
            "caseName": case["caseName"],
            "source": case["source"],
            "json": filename,
            "actions": len(case["initialization"]) + len(case["steps"]) + len(case["cleanup"]),
            "warnings": len(case["warnings"]),
        })

    manifest = {
        "totalCases": len(cases),
        "totalWarnings": sum(len(case["warnings"]) for case in cases),
        "cases": manifest_cases,
    }
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return manifest


def extract_first_case(workbook_path, sheet_name="3.1 驾驶模式"):
    """保留旧调用接口。"""
    return copy.deepcopy(extract_workbook(workbook_path, [sheet_name])[0])


def main():
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="批量解析Excel中的HIL测试用例")
    parser.add_argument("input", nargs="?", default=str(project_root / "用例.xlsx"))
    parser.add_argument("--sheet", action="append", help="只解析指定工作表，可重复")
    parser.add_argument(
        "--output-dir",
        default=str(Path(__file__).with_name("generated_cases") / "json"),
    )
    args = parser.parse_args()

    cases = extract_workbook(args.input, args.sheet)
    manifest = write_cases(cases, Path(args.output_dir))
    print(f"解析完成：{manifest['totalCases']}条用例")
    print(f"警告总数：{manifest['totalWarnings']}")
    print(f"输出目录：{Path(args.output_dir).resolve()}")


if __name__ == "__main__":
    main()
